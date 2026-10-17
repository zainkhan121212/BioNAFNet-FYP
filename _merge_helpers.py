"""
Shared formatting helpers for the merged FYP Master Report.
Uses Word built-in Heading styles (for TOC) and Caption style with SEQ fields (for LOF/LOT).
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# ── Style constants (from Handbook) ──
FONT_BODY = "Times New Roman"
FONT_HEADING2 = "Arial"
FONT_HEADING3 = "Arial"
SIZE_BODY = Pt(12)
SIZE_H1 = Pt(20)
SIZE_H2 = Pt(14)
SIZE_H3 = Pt(13)
LINE_SPACING = 1.5
SPACE_BEFORE = Pt(12)
SPACE_AFTER = Pt(6)


def _clear_theme_formatting(style):
    """Remove theme color and font references so explicit settings take effect."""
    rPr = style.element.get_or_add_rPr()
    color_elem = rPr.find(qn('w:color'))
    if color_elem is not None:
        for a in ['w:themeColor', 'w:themeShade', 'w:themeTint']:
            k = qn(a)
            if k in color_elem.attrib:
                del color_elem.attrib[k]
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        for a in ['w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme']:
            k = qn(a)
            if k in rFonts.attrib:
                del rFonts.attrib[k]


def setup_styles(doc):
    """Configure document styles per handbook guidelines."""
    # Normal body style
    style = doc.styles['Normal']
    font = style.font
    font.name = FONT_BODY
    font.size = SIZE_BODY
    font.color.rgb = RGBColor(0, 0, 0)
    pf = style.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_before = SPACE_BEFORE
    pf.space_after = SPACE_AFTER
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Heading 1: Times New Roman, Bold, 20pt, Right align (Chapter headings)
    h1 = doc.styles['Heading 1']
    _clear_theme_formatting(h1)
    h1.font.name = FONT_BODY
    h1.font.size = SIZE_H1
    h1.font.bold = True
    h1.font.italic = False
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h1.paragraph_format.space_before = Pt(36)
    h1.paragraph_format.space_after = Pt(18)
    h1.paragraph_format.line_spacing = LINE_SPACING

    # Heading 2: Arial, Bold, Italic, 14pt, Left align (Section headings)
    h2 = doc.styles['Heading 2']
    _clear_theme_formatting(h2)
    h2.font.name = FONT_HEADING2
    h2.font.size = SIZE_H2
    h2.font.bold = True
    h2.font.italic = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(12)
    h2.paragraph_format.line_spacing = LINE_SPACING

    # Heading 3: Arial, Bold, 13pt, Left align (Sub-section headings)
    h3 = doc.styles['Heading 3']
    _clear_theme_formatting(h3)
    h3.font.name = FONT_HEADING3
    h3.font.size = SIZE_H3
    h3.font.bold = True
    h3.font.italic = False
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(8)
    h3.paragraph_format.line_spacing = LINE_SPACING

    # Caption style (for figure/table captions — enables LOF/LOT in Word)
    try:
        cap = doc.styles['Caption']
    except KeyError:
        from docx.enum.style import WD_STYLE_TYPE
        cap = doc.styles.add_style('Caption', WD_STYLE_TYPE.PARAGRAPH)
    cap.font.name = FONT_BODY
    cap.font.size = SIZE_BODY
    cap.font.bold = True
    cap.font.italic = False
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(2.54)


# ── Heading functions (use built-in styles for TOC support) ──

def add_heading1(doc, text):
    """Chapter heading: built-in Heading 1 style."""
    p = doc.add_heading(text, level=1)
    return p


def add_heading2(doc, text):
    """Topic heading: built-in Heading 2 style."""
    p = doc.add_heading(text, level=2)
    return p


def add_heading3(doc, text):
    """Sub heading: built-in Heading 3 style."""
    p = doc.add_heading(text, level=3)
    return p


# ── Paragraph helpers (unchanged formatting) ──

def add_para(doc, text, bold=False, italic=False, alignment=None):
    """Add a body paragraph with handbook formatting."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.font.size = SIZE_BODY
    run.bold = bold
    run.italic = italic
    p.paragraph_format.line_spacing = LINE_SPACING
    p.paragraph_format.space_before = SPACE_BEFORE
    p.paragraph_format.space_after = SPACE_AFTER
    if alignment:
        p.alignment = alignment
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_center(doc, text, bold=False, size=None):
    """Add centered text (for title pages, etc.)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.font.size = size if size else SIZE_BODY
    run.bold = bold
    p.paragraph_format.line_spacing = LINE_SPACING
    return p


def add_list_numbered(doc, text):
    """Add a numbered list item."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.font.size = SIZE_BODY
    p.paragraph_format.line_spacing = LINE_SPACING
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.27)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_list_bullet(doc, text):
    """Add a bullet list item."""
    p = doc.add_paragraph()
    run = p.add_run("\u2022  " + text)
    run.font.name = FONT_BODY
    run.font.size = SIZE_BODY
    p.paragraph_format.line_spacing = LINE_SPACING
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.27)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


# ── Caption helpers with SEQ fields (enables LOF/LOT in Word) ──

def _add_seq_field(paragraph, label, display_value):
    """Insert a SEQ field for auto-numbering figures/tables in Word."""
    run1 = paragraph.add_run()
    fc1 = OxmlElement('w:fldChar')
    fc1.set(qn('w:fldCharType'), 'begin')
    run1._r.append(fc1)

    run2 = paragraph.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = f' SEQ {label} \\* ARABIC '
    run2._r.append(instr)

    run3 = paragraph.add_run()
    fc2 = OxmlElement('w:fldChar')
    fc2.set(qn('w:fldCharType'), 'separate')
    run3._r.append(fc2)

    run4 = paragraph.add_run(display_value)
    run4.bold = True
    run4.font.name = FONT_BODY
    run4.font.size = SIZE_BODY

    run5 = paragraph.add_run()
    fc3 = OxmlElement('w:fldChar')
    fc3.set(qn('w:fldCharType'), 'end')
    run5._r.append(fc3)


def add_table_caption(doc, number, text):
    """Table caption ABOVE table with SEQ field (for LOT support)."""
    cap = doc.add_paragraph()
    cap.style = doc.styles['Caption']
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = cap.add_run("Table ")
    r1.bold = True
    r1.font.name = FONT_BODY
    r1.font.size = SIZE_BODY
    _add_seq_field(cap, "Table", str(number))
    r2 = cap.add_run(f": {text}")
    r2.bold = True
    r2.font.name = FONT_BODY
    r2.font.size = SIZE_BODY
    cap.paragraph_format.space_before = Pt(12)
    cap.paragraph_format.space_after = Pt(6)
    return cap


def add_figure_placeholder(doc, number, description, image_path=None):
    """Figure with caption BELOW. Inserts actual image if image_path provided."""
    if image_path and os.path.exists(image_path):
        doc.add_picture(image_path, width=Inches(5.5))
        pic_para = doc.paragraphs[-1]
        pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_para.paragraph_format.space_before = Pt(12)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[INSERT FIGURE {number} HERE]")
        run.bold = True
        run.italic = True
        run.font.name = FONT_BODY
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(128, 128, 128)
        p.paragraph_format.space_before = Pt(12)

    # Caption with SEQ field
    cap = doc.add_paragraph()
    cap.style = doc.styles['Caption']
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = cap.add_run("Figure ")
    r1.bold = True
    r1.font.name = FONT_BODY
    r1.font.size = SIZE_BODY
    _add_seq_field(cap, "Figure", str(number))
    r2 = cap.add_run(f": {description}")
    r2.bold = True
    r2.font.name = FONT_BODY
    r2.font.size = SIZE_BODY
    cap.paragraph_format.space_after = Pt(12)
    return cap


# ── TOC / LOF / LOT field code helpers ──

def _add_field_code(paragraph, field_code, placeholder_text):
    """Insert a Word field code into a paragraph."""
    run1 = paragraph.add_run()
    fc1 = OxmlElement('w:fldChar')
    fc1.set(qn('w:fldCharType'), 'begin')
    run1._r.append(fc1)

    run2 = paragraph.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = field_code
    run2._r.append(instr)

    run3 = paragraph.add_run()
    fc2 = OxmlElement('w:fldChar')
    fc2.set(qn('w:fldCharType'), 'separate')
    run3._r.append(fc2)

    run4 = paragraph.add_run(placeholder_text)
    run4.font.color.rgb = RGBColor(128, 128, 128)
    run4.font.size = SIZE_BODY
    run4.font.name = FONT_BODY

    run5 = paragraph.add_run()
    fc3 = OxmlElement('w:fldChar')
    fc3.set(qn('w:fldCharType'), 'end')
    run5._r.append(fc3)


def add_toc_field(doc):
    """Insert a Table of Contents field (user updates in Word)."""
    p = doc.add_paragraph()
    _add_field_code(p, ' TOC \\o "1-3" \\h \\z \\u ',
                    "Right-click here and select Update Field to generate Table of Contents")
    return p


def add_lof_field(doc):
    """Insert a List of Figures field (user updates in Word)."""
    p = doc.add_paragraph()
    _add_field_code(p, ' TOC \\h \\z \\c "Figure" ',
                    "Right-click here and select Update Field to generate List of Figures")
    return p


def add_lot_field(doc):
    """Insert a List of Tables field (user updates in Word)."""
    p = doc.add_paragraph()
    _add_field_code(p, ' TOC \\h \\z \\c "Table" ',
                    "Right-click here and select Update Field to generate List of Tables")
    return p


# ── Table helpers (unchanged) ──

def make_table(doc, headers, rows):
    """Create a formatted table with header row."""
    ncols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(str(h))
        run.bold = True
        run.font.name = FONT_BODY
        run.font.size = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade_cell(cell, "D9E2F3")

    for r_idx, row_data in enumerate(rows):
        row = table.rows[1 + r_idx]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.name = FONT_BODY
            run.font.size = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _set_table_borders(table)
    return table


def _shade_cell(cell, color_hex):
    """Apply background shading to a cell."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear'
    })
    shading.append(shd)


def _set_table_borders(table):
    """Set thin borders on all table cells."""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.makeelement(qn('w:tblPr'), {})
    borders = tblPr.makeelement(qn('w:tblBorders'), {})
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = borders.makeelement(qn(f'w:{edge}'), {
            qn('w:val'): 'single',
            qn('w:sz'): '4',
            qn('w:space'): '0',
            qn('w:color'): '000000'
        })
        borders.append(el)
    tblPr.append(borders)


def add_page_break(doc):
    """Insert a page break."""
    doc.add_page_break()
