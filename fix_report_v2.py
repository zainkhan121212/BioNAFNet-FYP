#!/usr/bin/env python3
"""
fix_report_v2.py - Second pass fixes for BioNAFNet Master Report

1. Remove ALL ensemble/synergy/fusion content (paragraphs, table rows, references)
2. Fix table numbering (start from Gantt Chart = Table 1, remove abbreviations as table)
3. Fill section 1.7 Significance with relevant content
4. Rebuild List of Tables after changes
5. Renumber all Table X: captions sequentially
"""

import sys, io, os, re, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC_PATH = r'F:\low_light_enhancement\eye adaption\BioNAFNet\BioNAFNet_Master_Reportnew_FIXED.docx'


def remove_ensemble_content(doc):
    """Remove all ensemble/synergy/fusion content from the document."""
    print("\n[STEP 1] Removing all ensemble/synergy/fusion content...")
    
    body = doc.element.body
    removed_paras = 0
    removed_rows = 0
    
    # --- Remove ensemble-related paragraphs ---
    ensemble_terms = ['ensemble', 'synergy', 'fusion', 'T60/S40', 'T70/S30', 'T50/S50', 
                      'group model', 'Ens T60', 'Ens 50/50', 'BioNAFNet-Ens']
    
    paras_to_remove = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip().lower()
        if not text:
            continue
        for term in ensemble_terms:
            if term.lower() in text:
                paras_to_remove.append(i)
                break
    
    # Remove in reverse order to avoid index shifting
    for i in reversed(paras_to_remove):
        p = doc.paragraphs[i]
        print(f"  Removing para {i}: {p.text.strip()[:80]}...")
        body.remove(p._element)
        removed_paras += 1
    
    print(f"  Removed {removed_paras} paragraphs")
    
    # --- Remove ensemble-related rows from tables ---
    for t_idx, table in enumerate(doc.tables):
        rows_to_remove = []
        for r_idx, row in enumerate(table.rows):
            if r_idx == 0:  # Skip header row
                continue
            row_text = ' '.join(c.text.strip() for c in row.cells).lower()
            for term in ensemble_terms:
                if term.lower() in row_text:
                    rows_to_remove.append(r_idx)
                    break
        
        # Remove rows in reverse order
        for r_idx in reversed(rows_to_remove):
            row = table.rows[r_idx]
            row_display = ' | '.join(c.text.strip()[:30] for c in row.cells)
            print(f"  Removing Table {t_idx}, Row {r_idx}: {row_display}")
            # Remove the row's XML element from the table
            table._tbl.remove(row._tr)
            removed_rows += 1
    
    print(f"  Removed {removed_rows} table rows")
    
    # --- Clean up RQ4 reference ---
    # RQ4 references fusion, need to rewrite it or remove it
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if 'RQ4' in p.text:
            # Remove this research question
            print(f"  Removing RQ4 paragraph: {p.text.strip()[:80]}")
            body.remove(p._element)
            removed_paras += 1
            break
    
    # Renumber remaining RQs (RQ1-RQ3 stay, anything after RQ4 renumbered)
    # No action needed since RQ4 was the last one
    
    return removed_paras + removed_rows


def fix_abbreviations_caption(doc):
    """Remove Table numbering from List of Abbreviations caption."""
    print("\n[STEP 2] Fixing abbreviations - removing table number...")
    
    for i, p in enumerate(doc.paragraphs):
        if p.style.name == 'Caption' and 'List of Abbreviations' in p.text:
            # Change from "Table 1: List of Abbreviations" to just have no table number
            # Actually, change the caption style to Normal since it shouldn't be in the
            # table numbering sequence
            old_text = p.text.strip()
            # Clear all runs and set new text
            for run in p.runs:
                run.text = ''
            if p.runs:
                p.runs[0].text = 'List of Abbreviations'
            # Change style to Normal so it doesn't appear in List of Tables
            p.style = doc.styles['Normal']
            # Make it bold and centered
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(14)
            p.paragraph_format.alignment = 1  # Center
            print(f"  Changed: '{old_text}' -> 'List of Abbreviations' (Normal style)")
            break


def renumber_tables(doc):
    """Renumber all Table captions sequentially starting from 1."""
    print("\n[STEP 3] Renumbering table captions...")
    
    table_num = 0
    old_to_new = {}
    
    # First pass: collect and renumber captions
    for i, p in enumerate(doc.paragraphs):
        if p.style.name == 'Caption' and p.text.strip().startswith('Table'):
            text = p.text.strip()
            # Extract old table number
            m = re.match(r'^Table\s+(\d+):\s*(.*)', text)
            if m:
                old_num = int(m.group(1))
                title = m.group(2)
                table_num += 1
                old_to_new[old_num] = table_num
                
                new_text = f'Table {table_num}: {title}'
                
                # Update the paragraph text
                for run in p.runs:
                    run.text = ''
                if p.runs:
                    p.runs[0].text = new_text
                else:
                    p.add_run(new_text)
                
                if old_num != table_num:
                    print(f"  Table {old_num} -> Table {table_num}: {title[:60]}")
    
    print(f"  Total tables renumbered: {table_num}")
    print(f"  Mapping: {old_to_new}")
    return old_to_new


def fill_section_17(doc):
    """Fill section 1.7 Significance with relevant content."""
    print("\n[STEP 4] Filling section 1.7 Significance...")
    
    body = doc.element.body
    
    # Find section 1.7
    target_idx = None
    next_heading_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == '1.7 Significance':
            target_idx = i
        elif target_idx is not None and p.style.name.startswith('Heading') and p.text.strip():
            next_heading_idx = i
            break
    
    if target_idx is None:
        print("  ERROR: Could not find section 1.7!")
        return
    
    print(f"  Found 1.7 at para {target_idx}, next heading at para {next_heading_idx}")
    
    # Content for section 1.7 Significance
    significance_content = [
        "This project holds significant value across both academic research and practical deployment of low-light image enhancement systems:",
        "Academic Significance: BioNAFNet is the first architecture to integrate seven distinct biological visual processing mechanisms into a single deep learning framework for low-light image enhancement. Unlike previous approaches that rely purely on mathematical formulations, our biologically-inspired modules (RLIN, FAEB, ARFP, PRCM, SPDP, TAM, WFPL) replicate actual retinal processing pathways, offering a novel paradigm for image restoration research. This work contributes new evidence that biological principles can significantly improve computational image processing.",
        "Practical Significance: The knowledge distillation pipeline compresses the 154M-parameter Teacher model into a 24M-parameter Student model with 6.4x reduction in size and 3.9x improvement in inference speed, making real-time enhancement feasible. The Student model achieves 20.5 FPS, enabling live webcam processing and mobile deployment. This addresses the critical industry gap between enhancement quality and computational efficiency.",
        "Industrial Significance: The accompanying NeonEnhancer Pro desktop application demonstrates a complete software solution integrating twelve enhancement models with batch processing, live preview, augmentation, cropping, and resizing capabilities. This provides a practical tool for security camera operators, mobile photographers, and content creators who regularly encounter low-light imaging challenges.",
        "Furthermore, the project establishes comprehensive benchmarks across three standard datasets (LOLv1, LOLv2 Real, LOLv2 Synthetic) with detailed comparisons against state-of-the-art methods, contributing valuable reference data for future research in the field.",
    ]
    
    # Insert content after the heading
    insert_ref = doc.paragraphs[next_heading_idx]._element
    
    # Insert in reverse order
    for text in reversed(significance_content):
        new_p = OxmlElement('w:p')
        
        # Paragraph properties - Normal style
        pPr = OxmlElement('w:pPr')
        pStyle = OxmlElement('w:pStyle')
        pStyle.set(qn('w:val'), 'Normal')
        pPr.append(pStyle)
        new_p.append(pPr)
        
        # Text run
        run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rPr.append(rFonts)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '24')  # 12pt
        rPr.append(sz)
        run.append(rPr)
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        run.append(t)
        new_p.append(run)
        
        insert_ref.addprevious(new_p)
    
    print(f"  Added {len(significance_content)} paragraphs to section 1.7")


def rebuild_lot_lof(doc):
    """Rebuild List of Tables and List of Figures after changes."""
    print("\n[STEP 5] Rebuilding List of Tables and List of Figures...")
    
    paras = doc.paragraphs
    body = doc.element.body
    
    # Collect current figure and table captions from body
    figure_captions = []
    table_captions = []
    
    # Find where body starts
    body_start = 0
    for i, p in enumerate(paras):
        if p.style.name == 'Heading 1' and p.text.strip().startswith('Chapter'):
            body_start = i
            break
    
    for i, p in enumerate(paras):
        if i < body_start:
            continue
        if p.style.name == 'Caption':
            text = p.text.strip()
            if text.startswith('Figure') and ':' in text:
                figure_captions.append(text)
            elif text.startswith('Table') and ':' in text:
                if 'List of Abbreviations' not in text and 'Table of Contents' not in text:
                    table_captions.append(text)
    
    print(f"  Found {len(table_captions)} table captions")
    print(f"  Found {len(figure_captions)} figure captions")
    
    # Find and remove existing LOT section
    lot_start = None
    lot_end = None
    lof_start = None
    lof_end = None
    
    for i, p in enumerate(paras):
        t = p.text.strip()
        if t == 'List of Tables':
            lot_start = i
        elif t == 'List of Figures' and lot_start is not None:
            lot_end = i - 1
            lof_start = i
        elif t == 'List of Abbreviations' and lof_start is not None:
            lof_end = i - 1
            break
    
    if lot_start is None or lof_start is None:
        print("  WARNING: Could not find LOT/LOF sections!")
        return
    
    # Find anchor (List of Abbreviations paragraph)
    anchor = None
    for p in paras:
        if p.text.strip() == 'List of Abbreviations':
            anchor = p._element
            break
    
    if anchor is None:
        print("  WARNING: Could not find anchor!")
        return
    
    # Remove everything from lot_start to lof_end
    elements_to_remove = []
    for i in range(lot_start, lof_end + 1):
        elements_to_remove.append(paras[i]._element)
    
    for elem in elements_to_remove:
        try:
            body.remove(elem)
        except:
            pass
    
    print(f"  Removed {len(elements_to_remove)} old LOT/LOF entries")
    
    # Create LOT header
    def make_header(text):
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), 'center')
        pPr.append(jc)
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '240')
        spacing.set(qn('w:after'), '120')
        pPr.append(spacing)
        p.append(pPr)
        run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rPr.append(rFonts)
        b = OxmlElement('w:b')
        rPr.append(b)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '28')
        rPr.append(sz)
        run.append(rPr)
        t_elem = OxmlElement('w:t')
        t_elem.text = text
        run.append(t_elem)
        p.append(run)
        return p
    
    def make_entry(text):
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '60')
        spacing.set(qn('w:after'), '60')
        spacing.set(qn('w:line'), '360')
        spacing.set(qn('w:lineRule'), 'auto')
        pPr.append(spacing)
        p.append(pPr)
        run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rPr.append(rFonts)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '24')
        rPr.append(sz)
        run.append(rPr)
        t_elem = OxmlElement('w:t')
        t_elem.set(qn('xml:space'), 'preserve')
        t_elem.text = text
        run.append(t_elem)
        p.append(run)
        return p
    
    def make_blank():
        return OxmlElement('w:p')
    
    # Re-find anchor after removal
    paras = doc.paragraphs
    anchor = None
    for p in paras:
        if p.text.strip() == 'List of Abbreviations':
            anchor = p._element
            break
    
    # Insert LOT
    anchor.addprevious(make_header('List of Tables'))
    anchor.addprevious(make_blank())
    for cap in table_captions:
        anchor.addprevious(make_entry(cap))
    anchor.addprevious(make_blank())
    anchor.addprevious(make_blank())
    
    # Insert LOF
    anchor.addprevious(make_header('List of Figures'))
    anchor.addprevious(make_blank())
    for cap in figure_captions:
        anchor.addprevious(make_entry(cap))
    anchor.addprevious(make_blank())
    anchor.addprevious(make_blank())
    
    print(f"  Created LOT with {len(table_captions)} entries")
    print(f"  Created LOF with {len(figure_captions)} entries")


def rebuild_toc(doc):
    """Rebuild the TOC to reflect heading changes."""
    print("\n[STEP 6] Rebuilding Table of Contents...")
    
    paras = doc.paragraphs
    body = doc.element.body
    
    # Find TOC region by checking XML style
    toc_start = None
    toc_end = None
    for i, p in enumerate(paras):
        pPr = p._element.find(qn('w:pPr'))
        if pPr is not None:
            pStyle = pPr.find(qn('w:pStyle'))
            if pStyle is not None and pStyle.get(qn('w:val'), '').startswith('toc'):
                if toc_start is None:
                    toc_start = i
                toc_end = i
    
    if toc_start is None:
        print("  WARNING: No TOC entries found!")
        return
    
    print(f"  Current TOC: paras {toc_start} to {toc_end}")
    
    # Get element after TOC
    after_toc = paras[toc_end]._element.getnext()
    
    # Remove old TOC
    for i in range(toc_start, toc_end + 1):
        try:
            body.remove(paras[i]._element)
        except:
            pass
    
    # Collect headings from body
    paras = doc.paragraphs  # Re-read after removal
    headings = []
    body_started = False
    for p in paras:
        if p.style.name == 'Heading 1' and p.text.strip().startswith('Chapter'):
            body_started = True
        if body_started and p.style.name.startswith('Heading') and p.text.strip():
            level = int(p.style.name.split()[-1])
            headings.append((level, p.text.strip()))
    
    # Also add Appendices and References
    for p in paras:
        if p.style.name.startswith('Heading') and p.text.strip() in ('Appendices', 'References'):
            level = int(p.style.name.split()[-1])
            headings.append((level, p.text.strip()))
        if p.style.name == 'Heading 2' and p.text.strip().startswith('Appendix'):
            headings.append((2, p.text.strip()))
    
    # Remove duplicates while preserving order
    seen = set()
    unique_headings = []
    for h in headings:
        if h not in seen:
            seen.add(h)
            unique_headings.append(h)
    headings = unique_headings
    
    print(f"  Found {len(headings)} headings")
    
    # Create TOC entries
    for level, text in headings:
        new_p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        pStyle = OxmlElement('w:pStyle')
        pStyle.set(qn('w:val'), f'toc {level}')
        pPr.append(pStyle)
        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'right')
        tab.set(qn('w:leader'), 'dot')
        tab.set(qn('w:pos'), '9360')
        tabs.append(tab)
        pPr.append(tabs)
        new_p.append(pPr)
        
        run = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        run.append(t)
        new_p.append(run)
        
        tab_run = OxmlElement('w:r')
        tab_char = OxmlElement('w:tab')
        tab_run.append(tab_char)
        new_p.append(tab_run)
        
        pn_run = OxmlElement('w:r')
        pn_t = OxmlElement('w:t')
        pn_t.text = '##'
        pn_run.append(pn_t)
        new_p.append(pn_run)
        
        if after_toc is not None:
            after_toc.addprevious(new_p)
        else:
            body.append(new_p)
    
    print(f"  Created {len(headings)} TOC entries")


def verify(doc):
    """Final verification."""
    print("\n[VERIFICATION]")
    
    # Check no ensemble content remains
    ensemble_found = False
    terms = ['ensemble', 'synergy', 'T60/S40', 'T70/S30', 'Ens T60', 'BioNAFNet-Ens']
    for p in doc.paragraphs:
        for term in terms:
            if term.lower() in p.text.lower():
                print(f"  WARN: ensemble content still found: {p.text.strip()[:80]}")
                ensemble_found = True
    
    for table in doc.tables:
        for row in table.rows:
            row_text = ' '.join(c.text for c in row.cells).lower()
            for term in terms:
                if term.lower() in row_text:
                    print(f"  WARN: ensemble in table: {' | '.join(c.text.strip()[:25] for c in row.cells)}")
                    ensemble_found = True
    
    if not ensemble_found:
        print("  Ensemble content: FULLY REMOVED")
    
    # Check section 1.7
    found_17 = False
    has_content = False
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == '1.7 Significance':
            found_17 = True
            continue
        if found_17 and p.style.name.startswith('Heading') and p.text.strip():
            break
        if found_17 and p.text.strip():
            has_content = True
    print(f"  Section 1.7 has content: {'YES' if has_content else 'NO'}")
    
    # Check table numbering
    table_nums = []
    for p in doc.paragraphs:
        if p.style.name == 'Caption' and p.text.strip().startswith('Table'):
            m = re.match(r'Table (\d+):', p.text.strip())
            if m:
                table_nums.append(int(m.group(1)))
    
    if table_nums:
        print(f"  Table numbering: starts at {table_nums[0]}, ends at {table_nums[-1]}, count={len(table_nums)}")
        # Check sequential
        expected = list(range(1, len(table_nums) + 1))
        if table_nums == expected:
            print("  Table numbering: SEQUENTIAL")
        else:
            print(f"  Table numbering: NOT sequential! {table_nums}")
    
    # Check abbreviations is not in table captions
    abbrev_in_tables = any('Abbreviation' in p.text for p in doc.paragraphs if p.style.name == 'Caption')
    print(f"  Abbreviations removed from table numbering: {'YES' if not abbrev_in_tables else 'NO'}")


def main():
    print("=" * 60)
    print("BioNAFNet Master Report - Fix V2")
    print("=" * 60)
    
    doc = Document(DOC_PATH)
    print(f"Loaded: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")
    
    remove_ensemble_content(doc)
    fix_abbreviations_caption(doc)
    renumber_tables(doc)
    fill_section_17(doc)
    rebuild_lot_lof(doc)
    rebuild_toc(doc)
    verify(doc)
    
    doc.save(DOC_PATH)
    fsize = os.path.getsize(DOC_PATH) / (1024 * 1024)
    print(f"\nSaved: {DOC_PATH} ({fsize:.1f} MB)")
    print("DONE!")


if __name__ == '__main__':
    main()
