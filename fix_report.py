#!/usr/bin/env python3
"""
fix_report.py - Comprehensive Fix for BioNAFNet Master Report

Fixes:
  1. Heading styles (body text wrongly styled as Heading 3 -> Normal)
  2. Caption styles (body text wrongly styled as Caption -> Normal)
  3. Apply formatting guidelines (fonts, sizes, alignment, spacing)
  4. Restore missing images from source document
  5. Rebuild clean Table of Contents
  6. Rebuild List of Figures and List of Tables
"""

import os
import re
import sys
import copy
import shutil
import tempfile
import io

# Force UTF-8 output
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, Inches, Emu, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement
from lxml import etree

# ============================================================
# Configuration
# ============================================================
NEW_DOC = r'F:\low_light_enhancement\eye adaption\BioNAFNet\BioNAFNet_Master_Reportnew.docx'
SRC_DOC = r'F:\low_light_enhancement\eye adaption\BioNAFNet\BioNAFNet_Master_Report.docx'
OUT_DOC = r'F:\low_light_enhancement\eye adaption\BioNAFNet\BioNAFNet_Master_Reportnew_FIXED.docx'

# ============================================================
# Utility Functions
# ============================================================

def has_image(para):
    """Check if a paragraph contains an inline image."""
    return len(para._element.findall('.//' + qn('a:blip'))) > 0


def get_image_rids(para):
    """Get all image relationship IDs from a paragraph."""
    blips = para._element.findall('.//' + qn('a:blip'))
    return [b.get(qn('r:embed')) for b in blips if b.get(qn('r:embed'))]


def is_real_heading3(text):
    """Determine if text is an actual Heading 3 (section heading) vs body text."""
    text = text.strip()
    if not text:
        return False
    # Section numbers: "4.2 Something" or "6.3.1 Something"
    if re.match(r'^\d+\.\d+', text):
        return True
    # Numbered list items: "1. Something", "2. Something"
    if re.match(r'^\d+\.\s', text):
        return False
    # Starts with non-alphanumeric (bullets, dashes, etc)
    if not text[0].isalnum():
        return False
    # Long text is body text
    if len(text) > 80:
        return False
    # Short alphanumeric text - likely a sub-heading title
    return True


def is_real_caption(text):
    """Determine if text is an actual figure/table caption."""
    text = text.strip()
    if not text:
        return False
    if re.match(r'^(Figure|Table)\s+\d+', text):
        return True
    return False


def normalize_fig_caption(text):
    """Extract a matching key from a figure caption for cross-doc matching."""
    text = text.strip()
    m = re.match(r'^(?:\s*)Figure\s+\d+[:.]\s*(.*)', text)
    if m:
        key = m.group(1).strip().lower()
        key = re.sub(r'\s+', ' ', key)[:50]
        return key
    return None


def set_font_xml(style, font_name):
    """Set font in style XML to handle all character sets."""
    rPr = style.element.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        style.element.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)


# ============================================================
# STEP 1: Fix Heading Styles
# ============================================================

def fix_heading_styles(doc):
    """Fix body text paragraphs wrongly styled as Heading 3."""
    print("\n[STEP 1] Fixing heading styles...")
    h3_fixed = 0
    empty_fixed = 0

    for i, p in enumerate(doc.paragraphs):
        if p.style.name == 'Heading 3':
            text = p.text.strip()
            if not text:
                p.style = doc.styles['Normal']
                empty_fixed += 1
            elif not is_real_heading3(text):
                p.style = doc.styles['Normal']
                h3_fixed += 1

    print(f"  Fixed {h3_fixed} body paragraphs (Heading 3 -> Normal)")
    print(f"  Fixed {empty_fixed} empty headings -> Normal")
    return h3_fixed + empty_fixed


# ============================================================
# STEP 2: Fix Caption Styles
# ============================================================

def fix_caption_styles(doc):
    """Fix body text paragraphs wrongly styled as Caption."""
    print("\n[STEP 2] Fixing caption styles...")
    fixed = 0

    for i, p in enumerate(doc.paragraphs):
        if p.style.name == 'Caption':
            text = p.text.strip()
            if text and not is_real_caption(text):
                p.style = doc.styles['Normal']
                fixed += 1

    # Also fix the incorrect "Table 63:" caption
    for i, p in enumerate(doc.paragraphs):
        if p.style.name == 'Caption' and 'Table 63' in p.text:
            # This should be Table 1: List of Abbreviations
            for run in p.runs:
                if 'Table 63' in run.text:
                    run.text = run.text.replace('Table 63', 'Table 1')
            print("  Fixed 'Table 63' -> 'Table 1' in abbreviations caption")

    print(f"  Fixed {fixed} body paragraphs (Caption -> Normal)")
    return fixed


# ============================================================
# STEP 3: Apply Formatting Guidelines
# ============================================================

def apply_formatting(doc):
    """Apply formatting guidelines to all document styles."""
    print("\n[STEP 3] Applying formatting guidelines...")

    # --- Heading 1: Times New Roman, Bold, 20pt, Right-aligned ---
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(20)
    h1.font.bold = True
    h1.font.italic = False
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_font_xml(h1, 'Times New Roman')
    print("  Heading 1: Times New Roman, Bold, 20pt, Right-aligned")

    # --- Heading 2: Arial, Bold Italic, 14pt ---
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Arial'
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.italic = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_font_xml(h2, 'Arial')
    print("  Heading 2: Arial, Bold Italic, 14pt")

    # --- Heading 3: Arial, Bold, 13pt ---
    h3 = doc.styles['Heading 3']
    h3.font.name = 'Arial'
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.font.italic = False
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    h3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_font_xml(h3, 'Arial')
    print("  Heading 3: Arial, Bold, 13pt")

    # --- Normal (body text): Times New Roman, 12pt, Justified ---
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.font.bold = False
    normal.font.italic = False
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_font_xml(normal, 'Times New Roman')
    print("  Normal: Times New Roman, 12pt, Justified, 1.5 spacing")

    # --- Caption: Times New Roman, 10pt, Italic, Centered ---
    cap = doc.styles['Caption']
    cap.font.name = 'Times New Roman'
    cap.font.size = Pt(10)
    cap.font.italic = True
    cap.font.bold = False
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_font_xml(cap, 'Times New Roman')
    print("  Caption: Times New Roman, 10pt, Italic, Centered")

    # --- TOC styles ---
    for level_num in [1, 2, 3]:
        style_name = f'toc {level_num}'
        try:
            toc_s = doc.styles[style_name]
            toc_s.font.name = 'Times New Roman'
            toc_s.font.size = Pt(12)
            toc_s.font.color.rgb = RGBColor(0, 0, 0)
            toc_s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            toc_s.paragraph_format.space_before = Pt(0)
            toc_s.paragraph_format.space_after = Pt(0)
            set_font_xml(toc_s, 'Times New Roman')
            # Indentation based on level
            indent = Inches(0.3 * (level_num - 1))
            toc_s.paragraph_format.left_indent = indent
        except KeyError:
            pass
    print("  TOC styles: Times New Roman, 12pt, proper indentation")

    # --- Also clear run-level overrides on heading paragraphs ---
    for p in doc.paragraphs:
        if p.style.name.startswith('Heading'):
            for run in p.runs:
                # Clear run-level font overrides so style takes effect
                rPr = run._element.find(qn('w:rPr'))
                if rPr is not None:
                    # Remove font name override
                    rFonts = rPr.find(qn('w:rFonts'))
                    if rFonts is not None:
                        rPr.remove(rFonts)
                    # Remove size override
                    sz = rPr.find(qn('w:sz'))
                    if sz is not None:
                        rPr.remove(sz)
                    szCs = rPr.find(qn('w:szCs'))
                    if szCs is not None:
                        rPr.remove(szCs)
                    # Remove color override
                    color = rPr.find(qn('w:color'))
                    if color is not None:
                        rPr.remove(color)

    print("  Cleared run-level font overrides on headings")


# ============================================================
# STEP 4: Restore Missing Images
# ============================================================

def restore_images(doc, src_doc, temp_dir):
    """Copy missing figure images from source document."""
    print("\n[STEP 4] Restoring missing images...")

    # Build source map: normalized caption key -> image info
    src_map = {}
    src_paras = src_doc.paragraphs

    for i, p in enumerate(src_paras):
        if p.style.name == 'Caption' and p.text.strip().startswith('Figure'):
            key = normalize_fig_caption(p.text)
            if not key:
                continue
            # Search for image in nearby paragraphs (prefer before caption)
            for offset in [-1, -2, -3, 0, 1, 2, 3]:
                j = i + offset
                if 0 <= j < len(src_paras) and has_image(src_paras[j]):
                    rids = get_image_rids(src_paras[j])
                    if rids:
                        src_map[key] = {'para_idx': j, 'rids': rids}
                        break

    print(f"  Source doc: mapped {len(src_map)} figure images")

    # Find figure captions in new doc missing their images
    paras = doc.paragraphs
    missing = []

    for i, p in enumerate(paras):
        if p.style.name == 'Caption' and p.text.strip().startswith('Figure'):
            found_img = False
            for offset in [-1, -2, -3, 0, 1, 2, 3]:
                j = i + offset
                if 0 <= j < len(paras) and has_image(paras[j]):
                    found_img = True
                    break
            if not found_img:
                key = normalize_fig_caption(p.text)
                if key and key in src_map:
                    missing.append((i, p.text.strip(), key))

    print(f"  New doc: {len(missing)} figures missing their images")

    if not missing:
        print("  No images to restore!")
        return 0

    # Process in REVERSE order to avoid index shifting
    restored = 0
    for para_idx, caption_text, caption_key in reversed(missing):
        try:
            src_info = src_map[caption_key]
            # Get the first image from the source paragraph
            rid = src_info['rids'][0]
            if rid not in src_doc.part.rels:
                print(f"  SKIP: Cannot find relationship {rid} in source")
                continue

            rel = src_doc.part.rels[rid]
            image_blob = rel.target_part.blob
            content_type = rel.target_part.content_type

            # Determine extension
            ext = '.png'
            if 'jpeg' in content_type or 'jpg' in content_type:
                ext = '.jpg'
            elif 'emf' in content_type:
                ext = '.emf'
            elif 'wmf' in content_type:
                ext = '.wmf'
            elif 'gif' in content_type:
                ext = '.gif'

            # Save image to temp
            temp_path = os.path.join(temp_dir, f'restore_{restored}{ext}')
            with open(temp_path, 'wb') as f:
                f.write(image_blob)

            # Add picture to document (appends to end)
            doc.add_picture(temp_path, width=Inches(5.5))

            # Get the newly added paragraph (last one in body before sectPr)
            body = doc.element.body
            all_p_elements = body.findall(qn('w:p'))
            pic_p_elem = all_p_elements[-1]

            # Center-align the picture paragraph
            pPr = pic_p_elem.find(qn('w:pPr'))
            if pPr is None:
                pPr = OxmlElement('w:pPr')
                pic_p_elem.insert(0, pPr)
            jc = pPr.find(qn('w:jc'))
            if jc is None:
                jc = OxmlElement('w:jc')
                pPr.append(jc)
            jc.set(qn('w:val'), 'center')

            # Move it before the caption paragraph
            # Re-index paragraphs since we might have shifted things
            caption_elem = doc.paragraphs[para_idx]._element
            caption_elem.addprevious(pic_p_elem)

            restored += 1
            cap_short = caption_text[:70]
            print(f"  Restored: {cap_short}...")

        except Exception as e:
            print(f"  ERROR restoring '{caption_text[:40]}': {e}")

    print(f"  Total restored: {restored} images")
    return restored


# ============================================================
# STEP 5: Rebuild Table of Contents
# ============================================================

def create_toc_entry(text, level, page_num="##"):
    """Create a single TOC entry paragraph XML element."""
    new_p = OxmlElement('w:p')

    # Paragraph properties
    pPr = OxmlElement('w:pPr')
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), f'toc {level}')
    pPr.append(pStyle)

    # Tab stop for right-aligned page number with dot leader
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), '9360')  # ~6.5 inches in twips
    tabs.append(tab)
    pPr.append(tabs)

    new_p.append(pPr)

    # Text run
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    new_p.append(run)

    # Tab character
    tab_run = OxmlElement('w:r')
    tab_char = OxmlElement('w:tab')
    tab_run.append(tab_char)
    new_p.append(tab_run)

    # Page number
    pn_run = OxmlElement('w:r')
    pn_t = OxmlElement('w:t')
    pn_t.text = str(page_num)
    pn_run.append(pn_t)
    new_p.append(pn_run)

    return new_p


def rebuild_toc(doc):
    """Rebuild the Table of Contents from actual headings."""
    print("\n[STEP 5] Rebuilding Table of Contents...")

    paras = doc.paragraphs

    # Find TOC region
    toc_start = None
    toc_end = None
    for i, p in enumerate(paras):
        if p.style.name in ('toc 1', 'toc 2', 'toc 3'):
            if toc_start is None:
                toc_start = i
            toc_end = i

    if toc_start is None:
        print("  WARNING: No TOC entries found!")
        return

    print(f"  Current TOC: paras {toc_start} to {toc_end} ({toc_end - toc_start + 1} entries)")

    # Save the old TOC entries for page number reference
    old_toc = {}
    for i in range(toc_start, toc_end + 1):
        p = paras[i]
        text = p.text.strip()
        if '\t' in text:
            parts = text.rsplit('\t', 1)
            old_toc[parts[0].strip()] = parts[1].strip()

    # Collect actual headings from body (after LOF/LOT/Abbreviations)
    # Find the first Heading 1 in body
    body_start = None
    for i, p in enumerate(paras):
        if i > toc_end and p.style.name == 'Heading 1' and p.text.strip():
            body_start = i
            break

    if body_start is None:
        print("  WARNING: No body headings found!")
        return

    headings = []
    for i, p in enumerate(paras):
        if i < body_start:
            continue
        if p.style.name.startswith('Heading'):
            text = p.text.strip()
            if text:
                level = int(p.style.name.split()[-1])
                headings.append((level, text))

    print(f"  Found {len(headings)} headings in document body")

    # Get the element after TOC for insertion reference
    body = doc.element.body
    after_toc_elem = paras[toc_end]._element.getnext()

    # Remove old TOC entries
    for i in range(toc_start, toc_end + 1):
        elem = paras[i]._element
        body.remove(elem)

    print(f"  Removed {toc_end - toc_start + 1} old TOC entries")

    # Create new TOC entries
    for level, text in headings:
        # Try to get old page number
        page = old_toc.get(text, '##')

        new_entry = create_toc_entry(text, level, page)

        # Insert before the after_toc element
        if after_toc_elem is not None:
            after_toc_elem.addprevious(new_entry)
        else:
            body.append(new_entry)

    print(f"  Created {len(headings)} clean TOC entries")


# ============================================================
# STEP 6: Rebuild List of Figures and List of Tables
# ============================================================

def create_lof_lot_entry(text):
    """Create a single LOF/LOT entry paragraph."""
    new_p = OxmlElement('w:p')

    # Paragraph properties - Normal style with tab stop
    pPr = OxmlElement('w:pPr')
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), 'Normal')
    pPr.append(pStyle)

    # Indentation
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '0')
    pPr.append(ind)

    # Spacing - tighter for list entries
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '60')
    spacing.set(qn('w:after'), '60')
    spacing.set(qn('w:line'), '360')  # 1.5 line spacing
    spacing.set(qn('w:lineRule'), 'auto')
    pPr.append(spacing)

    new_p.append(pPr)

    # Text run
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')  # 12pt
    rPr.append(sz)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    new_p.append(run)

    return new_p


def create_section_header(text):
    """Create a bold centered section header paragraph."""
    new_p = OxmlElement('w:p')

    pPr = OxmlElement('w:pPr')
    # Center alignment
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    pPr.append(jc)
    # Spacing
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '240')
    spacing.set(qn('w:after'), '120')
    pPr.append(spacing)
    new_p.append(pPr)

    # Bold text run
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    b = OxmlElement('w:b')
    rPr.append(b)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '28')  # 14pt
    rPr.append(sz)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    new_p.append(run)

    return new_p


def rebuild_lof_lot(doc):
    """Rebuild the List of Tables and List of Figures."""
    print("\n[STEP 6] Rebuilding List of Tables and List of Figures...")

    paras = doc.paragraphs
    body = doc.element.body

    # --- Find existing LOT/LOF regions ---
    lot_header = None
    lot_start = None
    lot_end = None
    lof_header = None
    lof_start = None
    lof_end = None

    for i, p in enumerate(paras):
        text = p.text.strip()
        if text == 'List of Tables':
            lot_header = i
        elif text == 'List of Figures':
            lof_header = i

    # Determine LOT entries (between "List of Tables" and "List of Figures")
    if lot_header is not None and lof_header is not None:
        lot_start = lot_header  # Starting from header itself
        lot_end = lof_header - 1  # Up to just before LOF header
        # Skip empty paras at end
        while lot_end > lot_start and not paras[lot_end].text.strip():
            lot_end -= 1

    # Determine LOF entries (between "List of Figures" and next section)
    if lof_header is not None:
        lof_start = lof_header
        # Find end: look for "List of Abbreviations" or "Abstract" or next section
        for i in range(lof_header + 1, min(lof_header + 100, len(paras))):
            text = paras[i].text.strip()
            if text in ('List of Abbreviations', 'Abstract') or paras[i].style.name.startswith('Heading'):
                lof_end = i - 1
                break
        else:
            lof_end = lof_header + 40  # Fallback
        # Skip empty paras at end
        while lof_end > lof_start and not paras[lof_end].text.strip():
            lof_end -= 1

    # --- Collect actual figure and table captions from body ---
    figure_captions = []
    table_captions = []

    # Find where body content starts (after front matter)
    body_content_start = 0
    for i, p in enumerate(paras):
        if p.style.name == 'Heading 1' and p.text.strip().startswith('Chapter'):
            body_content_start = i
            break

    for i, p in enumerate(paras):
        if i < body_content_start:
            continue
        if p.style.name == 'Caption':
            text = p.text.strip()
            if text.startswith('Figure'):
                figure_captions.append(text)
            elif text.startswith('Table'):
                table_captions.append(text)

    print(f"  Found {len(table_captions)} table captions in body")
    print(f"  Found {len(figure_captions)} figure captions in body")

    # --- Rebuild List of Tables ---
    if lot_start is not None and lot_end is not None:
        # Find element after LOT region
        after_lot_elem = paras[lot_end]._element.getnext()

        # Remove old entries (keeping 1 blank para between sections)
        for i in range(lot_start, lot_end + 1):
            try:
                body.remove(paras[i]._element)
            except:
                pass

        # Create new LOT section
        lot_entries = []
        lot_entries.append(create_section_header('List of Tables'))

        for caption in table_captions:
            lot_entries.append(create_lof_lot_entry(caption))

        # Insert
        for entry in reversed(lot_entries):
            if after_lot_elem is not None:
                after_lot_elem.addprevious(entry)
            else:
                body.append(entry)

        print(f"  Rebuilt List of Tables with {len(table_captions)} entries")
    else:
        print("  WARNING: Could not find List of Tables section")

    # Need to re-read paragraphs since we modified the document
    paras = doc.paragraphs

    # Re-find LOF header position
    lof_header = None
    for i, p in enumerate(paras):
        if p.text.strip() == 'List of Figures':
            lof_header = i
            break

    if lof_header is not None:
        # Find end of LOF section
        lof_end = None
        for i in range(lof_header + 1, min(lof_header + 100, len(paras))):
            text = paras[i].text.strip()
            if text in ('List of Abbreviations', 'Abstract') or paras[i].style.name.startswith('Heading'):
                lof_end = i - 1
                break
        if lof_end is None:
            lof_end = lof_header + 40

        while lof_end > lof_header and not paras[lof_end].text.strip():
            lof_end -= 1

        after_lof_elem = paras[lof_end]._element.getnext()

        # Remove old entries
        for i in range(lof_header, lof_end + 1):
            try:
                body.remove(paras[i]._element)
            except:
                pass

        # Create new LOF section
        lof_entries = []
        lof_entries.append(create_section_header('List of Figures'))

        for caption in figure_captions:
            lof_entries.append(create_lof_lot_entry(caption))

        # Insert
        for entry in reversed(lof_entries):
            if after_lof_elem is not None:
                after_lof_elem.addprevious(entry)
            else:
                body.append(entry)

        print(f"  Rebuilt List of Figures with {len(figure_captions)} entries")
    else:
        print("  WARNING: Could not find List of Figures section")


# ============================================================
# STEP 7: Verification
# ============================================================

def verify(doc):
    """Run verification checks on the fixed document."""
    print("\n[VERIFICATION] Running checks...")

    # Check 1: No body text with Heading 3 style
    bad_h3 = 0
    for p in doc.paragraphs:
        if p.style.name == 'Heading 3':
            if not is_real_heading3(p.text.strip()):
                bad_h3 += 1

    # Check 2: No body text with Caption style
    bad_cap = 0
    for p in doc.paragraphs:
        if p.style.name == 'Caption':
            if p.text.strip() and not is_real_caption(p.text.strip()):
                bad_cap += 1

    # Check 3: Count headings
    h1_count = sum(1 for p in doc.paragraphs if p.style.name == 'Heading 1' and p.text.strip())
    h2_count = sum(1 for p in doc.paragraphs if p.style.name == 'Heading 2' and p.text.strip())
    h3_count = sum(1 for p in doc.paragraphs if p.style.name == 'Heading 3' and p.text.strip())

    # Check 4: Count images
    img_count = sum(1 for p in doc.paragraphs if has_image(p))

    # Check 5: Count captions
    fig_caps = sum(1 for p in doc.paragraphs if p.style.name == 'Caption' and p.text.strip().startswith('Figure'))
    tbl_caps = sum(1 for p in doc.paragraphs if p.style.name == 'Caption' and p.text.strip().startswith('Table'))

    print(f"  Bad Heading 3 (body text): {bad_h3} {'PASS' if bad_h3 == 0 else 'FAIL'}")
    print(f"  Bad Caption (body text):   {bad_cap} {'PASS' if bad_cap == 0 else 'FAIL'}")
    print(f"  Heading 1 count: {h1_count}")
    print(f"  Heading 2 count: {h2_count}")
    print(f"  Heading 3 count: {h3_count}")
    print(f"  Paragraphs with images: {img_count}")
    print(f"  Figure captions: {fig_caps}")
    print(f"  Table captions:  {tbl_caps}")

    return bad_h3 == 0 and bad_cap == 0


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 60)
    print("BioNAFNet Master Report - Comprehensive Fix")
    print("=" * 60)

    # Create temp directory for image extraction
    temp_dir = os.path.join(os.path.dirname(NEW_DOC), '_temp_fix_images')
    os.makedirs(temp_dir, exist_ok=True)

    try:
        print(f"\nLoading documents...")
        doc = Document(NEW_DOC)
        src_doc = Document(SRC_DOC)
        print(f"  New doc:    {len(doc.paragraphs)} paragraphs")
        print(f"  Source doc: {len(src_doc.paragraphs)} paragraphs")

        # Execute all fix steps
        fix_heading_styles(doc)
        fix_caption_styles(doc)
        apply_formatting(doc)
        restore_images(doc, src_doc, temp_dir)
        rebuild_toc(doc)
        rebuild_lof_lot(doc)

        # Verify
        passed = verify(doc)

        # Save
        print(f"\n{'=' * 60}")
        print(f"Saving fixed document to:")
        print(f"  {OUT_DOC}")
        doc.save(OUT_DOC)
        file_size = os.path.getsize(OUT_DOC) / (1024 * 1024)
        print(f"  File size: {file_size:.1f} MB")
        print(f"\n{'=' * 60}")
        print(f"ALL DONE! {'All checks passed.' if passed else 'Some checks failed - review output above.'}")
        print(f"{'=' * 60}")
        print(f"\nNOTE: Open the document in Word, right-click the Table of")
        print(f"Contents, and select 'Update Field' > 'Update page numbers'")
        print(f"to get correct page numbers.")

    finally:
        # Clean up temp directory
        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir, ignore_errors=True)


if __name__ == '__main__':
    main()
