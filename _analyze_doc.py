import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn

doc = docx.Document(r'F:\low_light_enhancement\eye adaption\BioNAFNet\BioNAFNet_Master_Reportnew_FIXED.docx')

# 1. Verify NO ensemble anywhere
print('=== CHECK 1: No ensemble content ===')
terms = ['ensemble', 'synergy', 'fusion', 'T60/S40', 'T70/S30', 'BioNAFNet-Ens', 'Ens T60', 'Ens 50/50', 'group model']
found = False
for p in doc.paragraphs:
    for term in terms:
        if term.lower() in p.text.lower():
            print(f'  FOUND in para: {p.text.strip()[:80]}')
            found = True
for table in doc.tables:
    for row in table.rows:
        rt = ' '.join(c.text for c in row.cells).lower()
        for term in terms:
            if term.lower() in rt:
                print(f'  FOUND in table row: {" | ".join(c.text.strip()[:25] for c in row.cells)}')
                found = True
                break
print(f'  Result: {"FAIL - ensemble still exists" if found else "PASS - all removed"}')

# 2. Verify table numbering
print('\n=== CHECK 2: Table numbering ===')
for p in doc.paragraphs:
    if p.style.name == 'Caption' and p.text.strip().startswith('Table'):
        print(f'  {p.text.strip()[:80]}')
# Check no "List of Abbreviations" in Caption style
abbrev_cap = any('Abbreviation' in p.text for p in doc.paragraphs if p.style.name == 'Caption')
print(f'  Abbreviations in Caption style: {"YES (BAD)" if abbrev_cap else "NO (GOOD)"}')

# 3. Verify section 1.7
print('\n=== CHECK 3: Section 1.7 Significance ===')
in_17 = False
content_lines = 0
for p in doc.paragraphs:
    if p.text.strip() == '1.7 Significance':
        in_17 = True
        continue
    if in_17 and p.style.name.startswith('Heading') and p.text.strip():
        break
    if in_17 and p.text.strip():
        content_lines += 1
        print(f'  {p.text.strip()[:100]}...')
print(f'  Content paragraphs: {content_lines}')

# 4. Verify LOT
print('\n=== CHECK 4: List of Tables (first 5) ===')
in_lot = False
count = 0
for p in doc.paragraphs:
    if p.text.strip() == 'List of Tables':
        in_lot = True; continue
    if in_lot and p.text.strip() in ('List of Figures', 'List of Abbreviations'):
        in_lot = False; continue
    if in_lot and p.text.strip():
        count += 1
        if count <= 5: print(f'  {p.text.strip()[:80]}')
print(f'  Total LOT entries: {count}')

# 5. Verify LOF
print('\n=== CHECK 5: List of Figures (first 5) ===')
in_lof = False
count = 0
for p in doc.paragraphs:
    if p.text.strip() == 'List of Figures':
        in_lof = True; continue
    if in_lof and p.text.strip() in ('List of Abbreviations', 'Abstract'):
        in_lof = False; continue
    if in_lof and p.text.strip():
        count += 1
        if count <= 5: print(f'  {p.text.strip()[:80]}')
print(f'  Total LOF entries: {count}')

print('\n=== ALL CHECKS COMPLETE ===')
