"""
Master FYP Report Generator — Main Runner
Merges LLIE_FYP_Final_Report.docx and BioNAFNet_FYP_Final_Report_v2.docx
into a single comprehensive document with proper formatting.

Usage: py -3 _gen_master_run.py
Output: BioNAFNet_Master_Report.docx
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from _merge_helpers import setup_styles, add_page_break
from _gen_master_p1 import build_front_matter, build_chapter1, build_chapter2
from _gen_master_p2 import build_chapter3, build_chapter4
from _gen_master_p3 import build_chapter5, build_chapter6
from _gen_master_p4 import build_chapter7, build_chapter8, build_references, build_appendices

OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "BioNAFNet_Master_Report.docx")


def main():
    print("=" * 60)
    print("  BioNAFNet Master FYP Report Generator")
    print("  Merging two source documents into one master report")
    print("=" * 60)

    doc = Document()
    setup_styles(doc)

    print("[1/10] Building front matter (title, copyrights, abstract)...")
    build_front_matter(doc)

    print("[2/10] Building Chapter 1: Introduction...")
    build_chapter1(doc)

    print("[3/10] Building Chapter 2: Problem Definition...")
    build_chapter2(doc)

    print("[4/10] Building Chapter 3: Software Requirement Specification...")
    build_chapter3(doc)

    print("[5/10] Building Chapter 4: Methodology...")
    build_chapter4(doc)

    print("[6/10] Building Chapter 5: System Architecture...")
    build_chapter5(doc)

    print("[7/10] Building Chapter 6: Implementation and Testing...")
    build_chapter6(doc)

    print("[8/10] Building Chapter 7: Results and Discussion...")
    build_chapter7(doc)

    print("[9/10] Building Chapter 8: Conclusion and Future Work...")
    build_chapter8(doc)

    print("[10/10] Building References and Appendices...")
    build_references(doc)
    build_appendices(doc)

    print(f"\nSaving to: {OUTPUT}")
    doc.save(OUTPUT)

    # Count content
    n_para = len(doc.paragraphs)
    n_tables = len(doc.tables)
    print(f"\nDocument statistics:")
    print(f"  Paragraphs: {n_para}")
    print(f"  Tables:     {n_tables}")
    print(f"  Chapters:   8 + References + 6 Appendices")
    print(f"\nDone! Master report saved to: {OUTPUT}")


if __name__ == "__main__":
    main()
