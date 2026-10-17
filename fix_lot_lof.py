#!/usr/bin/env python3
"""
fix_lot_lof.py - Fix the List of Tables and List of Figures ordering
Runs on the already-fixed document to correct the LOT/LOF insertion order.
"""

import os
import sys
import io
import re

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC_PATH = r'F:\low_light_enhancement\eye adaption\BioNAFNet\BioNAFNet_Master_Reportnew_FIXED.docx'


def create_section_header(text):
    """Create a bold centered section header paragraph."""
    new_p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    pPr.append(jc)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '240')
    spacing.set(qn('w:after'), '120')
    pPr.append(spacing)
    new_p.append(pPr)

    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    b = OxmlElement('w:b')
    rPr.append(b)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '28')  # 14pt
    rPr.append(sz)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    new_p.append(run)
    return new_p


def create_entry(text):
    """Create a LOF/LOT entry paragraph."""
    new_p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')

    # Tab stop for alignment
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), '9360')
    tabs.append(tab)
    pPr.append(tabs)

    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '60')
    spacing.set(qn('w:after'), '60')
    spacing.set(qn('w:line'), '360')
    spacing.set(qn('w:lineRule'), 'auto')
    pPr.append(spacing)

    new_p.append(pPr)

    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')  # 12pt
    rPr.append(sz)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    new_p.append(run)

    return new_p


def main():
    print("Fixing List of Tables and List of Figures...")
    doc = Document(DOC_PATH)
    paras = doc.paragraphs
    body = doc.element.body

    # Find the entire LOT+LOF region
    # Region starts after the last TOC appendix entry and ends before "List of Abbreviations"
    region_start = None
    region_end = None
    
    # Find where TOC appendix entries end (look for "Appendix F" toc entry)
    for i, p in enumerate(paras):
        text = p.text.strip()
        if 'Appendix F' in text and i < 200:
            region_start = i + 1
            
    # Find "List of Abbreviations"
    for i, p in enumerate(paras):
        text = p.text.strip()
        if text == 'List of Abbreviations':
            region_end = i
            break
    
    if region_start is None or region_end is None:
        print(f"ERROR: Could not find region boundaries. start={region_start}, end={region_end}")
        return
    
    print(f"  LOT/LOF region: paras {region_start} to {region_end - 1}")
    
    # Get the element right after the region (the "List of Abbreviations" paragraph)
    anchor_elem = paras[region_end]._element
    
    # Remove ALL paragraphs in the LOT/LOF region
    elements_to_remove = []
    for i in range(region_start, region_end):
        elements_to_remove.append(paras[i]._element)
    
    for elem in elements_to_remove:
        try:
            body.remove(elem)
        except:
            pass
    
    print(f"  Removed {len(elements_to_remove)} old LOT/LOF paragraphs")
    
    # Collect actual figure and table captions from body content
    # Re-read paragraphs (they shifted after removal)
    paras = doc.paragraphs
    
    figure_captions = []
    table_captions = []
    
    for p in paras:
        if p.style.name == 'Caption':
            text = p.text.strip()
            if text.startswith('Figure') and ':' in text:
                figure_captions.append(text)
            elif text.startswith('Table') and ':' in text:
                # Skip "Table 1: Table of Contents" and "Table 1: List of Abbreviations"
                if 'Table of Contents' not in text and 'List of Abbreviations' not in text:
                    table_captions.append(text)
    
    print(f"  Found {len(table_captions)} table captions")
    print(f"  Found {len(figure_captions)} figure captions")
    
    # Re-find anchor after removal
    paras = doc.paragraphs
    anchor_elem = None
    for p in paras:
        if p.text.strip() == 'List of Abbreviations':
            anchor_elem = p._element
            break
    
    if anchor_elem is None:
        print("ERROR: Lost anchor element!")
        return
    
    # Build new LOT and LOF sections
    # Insert in FORWARD order, each one before the anchor
    
    # --- List of Tables ---
    lot_header = create_section_header('List of Tables')
    anchor_elem.addprevious(lot_header)
    
    # Add blank line
    blank_p = OxmlElement('w:p')
    anchor_elem.addprevious(blank_p)
    
    for caption in table_captions:
        entry = create_entry(caption)
        anchor_elem.addprevious(entry)
    
    # Add blank lines between sections
    blank_p2 = OxmlElement('w:p')
    anchor_elem.addprevious(blank_p2)
    blank_p3 = OxmlElement('w:p')
    anchor_elem.addprevious(blank_p3)
    
    # --- List of Figures ---
    lof_header = create_section_header('List of Figures')
    anchor_elem.addprevious(lof_header)
    
    blank_p4 = OxmlElement('w:p')
    anchor_elem.addprevious(blank_p4)
    
    for caption in figure_captions:
        entry = create_entry(caption)
        anchor_elem.addprevious(entry)
    
    # Add trailing blank line
    blank_p5 = OxmlElement('w:p')
    anchor_elem.addprevious(blank_p5)
    blank_p6 = OxmlElement('w:p')
    anchor_elem.addprevious(blank_p6)
    
    print(f"  Created List of Tables with {len(table_captions)} entries")
    print(f"  Created List of Figures with {len(figure_captions)} entries")
    
    # Fix the "Table 2: 2:" duplicate numbering
    for p in doc.paragraphs:
        if p.style.name == 'Caption' and 'Table 2: 2:' in p.text:
            for run in p.runs:
                if 'Table 2: 2:' in run.text:
                    run.text = run.text.replace('Table 2: 2:', 'Table 2:')
                    print("  Fixed 'Table 2: 2:' -> 'Table 2:' duplicate numbering")
    
    # Save
    doc.save(DOC_PATH)
    print(f"\n  Saved to: {DOC_PATH}")
    print("  DONE!")


if __name__ == '__main__':
    main()
